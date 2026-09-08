import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\admin\Downloads\CONCILIAÇÃO 0309.xlsx', data_only=True)
ws_os = wb['OS']

# Mapping of row ranges to store in sheet OS
# Planalto: R4 - R7
# Piraporinha: R9 - R15
# Mauá: R17 - R23
# Kennedy: R26 - R29
# Rudge Ramos: R32 - R42
# Santo André: R44 - R51
# Rei do Módulo: R53 - R58
# Jorge Beretta: R60 - R63
# Dom Pedro I: R65 - R70
# Jabaquara: R72 - R78

stores_map = [
    ('Planalto', 4, 7, 'st-06'),
    ('Piraporinha', 9, 15, 'st-05'),
    ('Mauá', 17, 23, '3a3dd7ce-fa8c-4aee-bac4-42f30fa6899f'),
    ('Kennedy', 26, 29, 'st-04'),
    ('Rudge Ramos', 32, 42, 'st-07'),
    ('Santo André', 44, 51, 'st-08'),
    ('Rei do Módulo', 53, 58, 'st-09'),
    ('Jorge Beretta', 60, 63, 'st-03'),
    ('Dom Pedro I', 65, 70, 'st-01'),
    ('Jabaquara', 72, 78, 'st-02')
]

print('=== DETALHAMENTO DE TODAS AS OSS DA SHEET OS ===\n')

total_patio = 0

for st_name, start_r, end_r, st_id in stores_map:
    print(f'[LOJA] {st_name} (Linhas {start_r} a {end_r}):')
    st_sum = 0
    for r in range(start_r + 2, end_r):
        c_os = ws_os.cell(r, 2).value
        c_dt = str(ws_os.cell(r, 3).value or '')[:10]
        c_val = ws_os.cell(r, 4).value
        c_pgto = str(ws_os.cell(r, 5).value or '')
        if c_os is not None:
            val = float(c_val or 0)
            st_sum += val
            print(f'   - OS #{str(c_os).strip():<6} | Data: {c_dt} | Saldo Patio: R$ {val:>8.2f} | Info Pgto: {c_pgto}')
    subtotal = float(ws_os.cell(end_r, 4).value or 0)
    print(f'   -> Subtotal Patio na Planilha: R$ {subtotal:.2f} (Soma calculada: R$ {st_sum:.2f})\n')
    total_patio += subtotal

print(f'TOTAL PATIO CONSOLIDADO: R$ {total_patio:.2f}')
